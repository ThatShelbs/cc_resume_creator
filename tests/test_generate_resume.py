"""Tests for the deterministic parts of the pipeline. No LLM calls."""

import os
import sys
from pathlib import Path

import docx
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import generate_resume as g  # noqa: E402


@pytest.fixture(autouse=True)
def _clear_warnings():
    g.WARNINGS.clear()
    yield
    g.WARNINGS.clear()


BANK = {
    "F001": {"id": "F001", "employer": "Acme", "text": "Saved $300k per year with MMM.", "metrics": ["$300k"]},
    "F002": {"id": "F002", "employer": "Globex", "text": "Built a churn model.", "metrics": []},
    "F003": {"id": "F003", "employer": "unassigned", "text": "Led a POC.", "metrics": []},
    "F004": {"id": "F004", "employer": "general", "text": "Python and SQL.", "metrics": []},
}


# --- citations -------------------------------------------------------------


def test_split_citation():
    assert g.split_citation("Did X [F001, f004]") == ("Did X.", ["F001", "F004"])
    assert g.split_citation("Did X.") == ("Did X.", [])


def test_validate_citations_drops_bad_provenance():
    bullets = {
        "Acme": [
            "Saved $300k per year. [F001]",
            "Unknown fact. [F999]",
            "Transplanted. [F002]",
            "Unassigned. [F003]",
            "Uncited bullet.",
            "General skill. [F004]",
        ]
    }
    cleaned, cites = g.validate_citations(bullets, BANK)
    assert cleaned["Acme"] == ["Saved $300k per year.", "Uncited bullet.", "General skill."]
    assert cites["Acme"] == [["F001"], [], ["F004"]]
    assert sum("dropped" in w for w in g.WARNINGS) == 3
    assert any("uncited" in w for w in g.WARNINGS)


def test_validate_citations_warns_on_number_not_in_cited_fact():
    g.validate_citations({"Acme": ["Saved $450k. [F001]"]}, BANK)
    assert any("450" in w for w in g.WARNINGS)


# --- skills ----------------------------------------------------------------


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("Python", "Python"),
        ("Team Leadership (8+ years managing direct reports)", "Team Leadership"),
        ("8+ years of leadership", None),
        ("Scaled revenue of a $1B+ brand", None),
        ("Led the analytics function", None),
        ("Uplift Modeling (CausalML)", "Uplift Modeling (CausalML)"),
        ("one two three four five six seven eight", None),
    ],
)
def test_clean_skill_label(raw, expected):
    assert g.clean_skill_label(raw) == expected


def test_split_skill_items_keeps_parenthesized_commas():
    items = g._split_skill_items("Python, Uplift Modeling (CausalML, PyLift) · SQL; Tableau | R")
    assert items == ["Python", "Uplift Modeling (CausalML, PyLift)", "SQL", "Tableau", "R"]


def test_extract_skills_categorized():
    draft = "## Core Competencies\n- **Tools:** Python | SQL\n- **Leadership:** Team Leadership, Vendor Management\n"
    assert g.extract_skills_from_draft(draft) == ["Python", "SQL", "Team Leadership", "Vendor Management"]


def test_extract_skills_flat_heading():
    draft = "## Summary\nText.\n\n## Skills\nPython, SQL, Forecasting\n\n## Experience\n- Did X\n"
    assert g.extract_skills_from_draft(draft) == ["Python", "SQL", "Forecasting"]


def test_merge_and_sort_skills_backfills_only_posting_terms():
    profile = {"tools": ["Python", "Tableau"], "skills": ["Forecasting"]}
    merged = g.merge_and_sort_skills(["sql", "Python (pandas)"], profile, "We need Python and Tableau.")
    assert merged == ["Python (pandas)", "sql", "Tableau"]


# --- guards ----------------------------------------------------------------


def test_deny_list_blocks_identities_not_collaboration():
    patterns = g.load_deny_patterns()
    ok = g.find_deny_violations("Partnered with software engineers.", {}, [], patterns)
    bad = g.find_deny_violations("Worked as a full-stack engineer.", {}, ["Figma"], patterns)
    assert ok == []
    assert len(bad) == 2


def test_fix_years_of_experience():
    src = "Data scientist with 10 years of experience."
    assert g.fix_years_of_experience("Leader with 15+ years of experience.", src) == (
        "Leader with 10+ years of experience."
    )
    assert g.fix_years_of_experience("Leader with 11 years of experience.", src).startswith("Leader with 11")


def test_fix_years_catches_loose_phrasing():
    src = "Data scientist with 10 years of experience."
    assert g.fix_years_of_experience("Over 15 years leading teams.", src) == "Over 10 years leading teams."
    assert g.fix_years_of_experience("Built 3 years of models.", src) == "Built 3 years of models."


def test_remove_em_dashes():
    assert g.remove_em_dashes("Built X — and Y") == "Built X, and Y"
    assert g.remove_em_dashes("Built X – and Y") == "Built X, and Y"
    assert g.remove_em_dashes("From 2019–2021") == "From 2019–2021"


def test_strip_cross_employer_mentions():
    out = g.strip_cross_employer_mentions(
        {"Allstate Insurance": ["Did X at Acme.", "Did Y."], "Acme Corp": ["Did Z."]}
    )
    assert out == {"Allstate Insurance": ["Did Y."], "Acme Corp": ["Did Z."]}
    assert len(g.WARNINGS) == 1


# --- fact bank + keyword coverage -------------------------------------------


def test_snap_employer():
    companies = ["Acme  Corp", "Globex"]
    assert g.snap_employer("acme corp", companies) == "Acme  Corp"
    assert g.snap_employer("General", companies) == "general"
    assert g.snap_employer("Initech", companies) is None


def test_load_fact_bank_flags_unknown_employer(tmp_path):
    path = tmp_path / "fact_bank.yaml"
    path.write_text(
        "facts:\n"
        "  - {id: F001, employer: acme, text: Did X}\n"
        "  - {id: F002, employer: Acmee, text: Did Y}\n",
        encoding="utf-8",
    )
    bank = g.load_fact_bank(["Acme"], path)
    assert bank["F001"]["employer"] == "Acme"
    assert bank["F002"]["employer"] == "unassigned"
    assert any("Acmee" in w and "F002" in w for w in g.WARNINGS)


def test_keyword_coverage():
    profile = {"tools": ["Python", "R", "Tableau"], "skills": ["Forecasting"]}
    bank = {"F001": {"tags": ["causal inference"]}}
    job = "Must know Python, R, and causal inference. Forecasting a plus."
    covered, missing = g.keyword_coverage(job, "Python expert in forecasting.", profile, bank)
    assert covered == ["Forecasting", "Python"]
    assert missing == ["causal inference"]  # "R" is under 2 chars and skipped


# --- file handling -----------------------------------------------------------


def test_find_latest_uses_mtime(tmp_path, monkeypatch):
    monkeypatch.setattr(g, "INPUT_DIR", tmp_path)
    older, newer = tmp_path / "in_job_zeta.docx", tmp_path / "in_job_alpha.docx"
    older.write_bytes(b"")
    newer.write_bytes(b"")
    os.utime(older, (1_000_000, 1_000_000))
    os.utime(newer, (2_000_000, 2_000_000))
    assert g.find_latest("in_job") == newer


def test_parse_prior_resume_roundtrip(tmp_path):
    d = docx.Document()
    for text, style in [
        ("EDUCATION", "Normal"),
        ("State University | Springfield\t2008 - 2012", "Normal"),
        ("B.S. Statistics", "Normal"),
        ("EXPERIENCE", "Normal"),
        ("Acme | Chicago, IL\t2015 - Present", "Normal"),
        ("Lead Data Scientist", "Normal"),
        ("Built a churn model.", "List Paragraph"),
        ("Saved $300k.", "List Paragraph"),
    ]:
        d.add_paragraph(text, style=style)
    path = tmp_path / "in_resume.docx"
    d.save(str(path))

    parsed = g.parse_prior_resume(path)
    assert parsed["education"][0]["institution"] == "State University"
    job = parsed["jobs"][0]
    assert (job["company"], job["location"], job["dates"], job["title"]) == (
        "Acme",
        "Chicago, IL",
        "2015 - Present",
        "Lead Data Scientist",
    )
    assert job["source_bullets"] == ["Built a churn model.", "Saved $300k."]


def test_build_docx_contains_content(tmp_path):
    data = {
        "first_name": "Ada",
        "last_name": "Lovelace",
        "email": "ada@example.com",
        "phone": "",
        "location": "London",
        "summary": "Analyst.",
        "education": [],
        "experience": [
            {"company": "Acme", "location": "", "dates": "2020", "title": "Lead", "bullets": ["Did X."]}
        ],
        "skills": ["Python"],
    }
    out = tmp_path / "out.docx"
    g.build_docx(data, out)
    text = g.read_docx_text(out)
    for expected in ("Ada Lovelace", "Analyst.", "Did X.", "Python"):
        assert expected in text


def test_contact_links_are_hyperlinks(tmp_path):
    data = {
        "first_name": "Ada",
        "last_name": "Lovelace",
        "email": "ada@example.com",
        "phone": "(555) 010-0199",
        "location": "London",
        "linkedin": "linkedin.com/in/ada",
        "summary": "Analyst.",
        "skills": [],
    }
    out = tmp_path / "out.docx"
    g.build_docx(data, out)
    d = docx.Document(str(out))
    targets = {r.target_ref for r in d.part.rels.values() if r.is_external}
    assert targets == {"mailto:ada@example.com", "https://linkedin.com/in/ada"}
    assert "ada@example.com | (555) 010-0199 | London | linkedin.com/in/ada" in g.read_docx_text(out)


def test_find_latest_job_accepts_text_formats(tmp_path, monkeypatch):
    monkeypatch.setattr(g, "INPUT_DIR", tmp_path)
    job = tmp_path / "in_job_example.txt"
    job.write_text("Director, Analytics\nPython required.", encoding="utf-8")
    (tmp_path / "in_job_notes.json").write_text("{}", encoding="utf-8")
    found = g.find_latest("in_job", g.JOB_POSTING_EXTS)
    assert found == job
    assert "Python required." in g.read_input_text(found)
    with pytest.raises(SystemExit):
        g.find_latest("in_profile")


def test_extract_citations_inline():
    text, ids = g.extract_citations("I cut churn 12% [F001] and led 4 analysts [F002, f001].")
    assert text == "I cut churn 12% and led 4 analysts."
    assert ids == ["F001", "F002"]


def test_validate_cover_letter():
    paragraphs = [
        "I saved $300k per year — with MMM. [F001]",
        "I did something unverifiable. [F999]",
        "I led a POC. [F003]",
        "Uncited middle.",
        "Uncited close.",
    ]
    result = g.validate_cover_letter(paragraphs, BANK, "Saved $300k per year with MMM.")
    assert result == [
        ("I saved $300k per year, with MMM.", ["F001"]),
        ("Uncited middle.", []),
        ("Uncited close.", []),
    ]
    assert sum("dropped" in w for w in g.WARNINGS) == 2
    uncited = [w for w in g.WARNINGS if "uncited cover letter" in w]
    assert len(uncited) == 1 and "paragraph 4" in uncited[0]


def test_build_cover_letter_docx(tmp_path):
    data = {"first_name": "Ada", "last_name": "Lovelace", "email": "ada@example.com", "phone": "", "location": ""}
    out = tmp_path / "cl.docx"
    g.build_cover_letter_docx(data, ["First paragraph.", "Second paragraph."], out)
    text = g.read_docx_text(out)
    for expected in ("Dear Hiring Team,", "First paragraph.", "Second paragraph.", "Sincerely,"):
        assert expected in text


def test_cover_letter_skill_loads_without_frontmatter():
    body = g.load_skill(g.COVER_LETTER_SKILL_PATH)
    assert not body.startswith("---")
    assert "<COVER_LETTER>" in body


def test_examples_parse(tmp_path):
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "examples"))
    import make_examples

    make_examples.main(tmp_path)
    contact = g.parse_applicant_info(tmp_path / "in_profile_example.docx")
    assert (contact["first_name"], contact["email"]) == ("Jordan", "jordan.rivera@example.com")
    jobs = g.parse_prior_resume(tmp_path / "in_resume_example.docx")["jobs"]
    companies = [j["company"] for j in jobs]
    assert companies == ["Northwind Traders", "Contoso Retail"]
    skills = g.parse_profile_skills(tmp_path / "in_profile_example.docx")
    assert skills == {"tools": ["Python", "SQL", "Tableau"], "skills": ["Forecasting", "Team Leadership", "Stakeholder Management"]}
    bank = g.load_fact_bank(companies, tmp_path / "fact_bank.example.yaml")
    assert len(bank) == 5 and not g.WARNINGS


def test_write_report(tmp_path):
    path = tmp_path / "r.md"
    g.write_report(
        path,
        {"Job posting": "in_job.docx"},
        "sonnet",
        "medium",
        ["something odd"],
        (["Python"], ["SQL"]),
        {"Acme": ["Saved $300k per year."]},
        {("Acme", "Saved $300k per year."): ["F001"]},
        BANK,
        "raw draft",
    )
    report = path.read_text(encoding="utf-8")
    assert "something odd" in report
    assert "`F001` Saved $300k per year with MMM." in report
    assert "Missing: SQL" in report
