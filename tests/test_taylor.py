"""Tests for Resume Taylor (the browser app's backend). No Claude calls: the
one test that runs a generation swaps in a fake pipeline script."""

import json
import shutil
import subprocess
import sys
import textwrap
import time
from pathlib import Path

import docx
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "examples"))

import generate_resume as g  # noqa: E402
from taylor import credentials as credentials_mod  # noqa: E402
from taylor import facts as facts_mod  # noqa: E402
from taylor.api import create_app  # noqa: E402
from taylor import profile_store, resume_ingest  # noqa: E402
from taylor.jobs import JobManager  # noqa: E402
from taylor.lint import lint_result  # noqa: E402
from taylor.paths import Paths  # noqa: E402
from taylor.projects import NotFound, ProjectError, ProjectStore, guess_company_role  # noqa: E402

import demo_data  # noqa: E402
import make_examples  # noqa: E402

JOB_TEXT = (ROOT / "examples" / "in_job_example.txt").read_text(encoding="utf-8")


def write_docx(paragraphs, path):
    d = docx.Document()
    for text, style in paragraphs:
        d.add_paragraph(text, style=style)
    d.save(str(path))


@pytest.fixture
def workspace(tmp_path):
    """A data root with the demo persona's inputs."""
    paths = Paths(tmp_path)
    paths.ensure()
    write_docx(demo_data.PROFILE, paths.input_dir / "in_profile.docx")
    write_docx(demo_data.RESUME, paths.input_dir / "in_resume_Jordan-Rivera.docx")
    paths.fact_bank_path.write_text(demo_data.FACT_BANK, encoding="utf-8")
    shutil.copy(ROOT / "tests" / "deny_fixture.txt", paths.deny_path)
    return paths


# ---------------------------------------------------------------------------
# Profile
# ---------------------------------------------------------------------------


def test_profile_round_trip(workspace, tmp_path):
    profile = profile_store.read_profile(workspace.input_dir / "in_profile.docx")
    assert profile.contact.name == "Jordan Rivera"
    assert profile.contact.linkedin.startswith("linkedin.com")
    assert [s.title for s in profile.sections][:2] == ["Leadership", "Projects"]

    profile.sections.append(profile_store.Section(title="Certifications", items=["  Spaced\ttext\nline  "]))
    out = tmp_path / "saved.docx"
    backup_dir = tmp_path / "archive"
    out.write_bytes(b"old")  # something to back up
    profile_store.save_profile(profile, out, backup_dir)

    assert list(backup_dir.iterdir()), "the previous file is backed up before overwriting"
    contact = g.parse_applicant_info(out)
    assert (contact["first_name"], contact["last_name"], contact["email"]) == ("Jordan", "Rivera", "jordan.rivera@example.com")
    skills = g.parse_profile_skills(out)
    assert "Python" in skills["tools"] and "Forecasting" in skills["skills"]
    again = profile_store.read_profile(out)
    assert again.sections[-1].items == ["Spaced text line"], "items are single clean paragraphs"


def test_profile_validation_rejects_missing_email():
    bad = profile_store.Profile(contact=profile_store.Contact(name="Jordan Rivera"))
    with pytest.raises(profile_store.ProfileError):
        profile_store.validate_profile(bad)


# ---------------------------------------------------------------------------
# Resume ingestion
# ---------------------------------------------------------------------------

PDF_LIKE = textwrap.dedent("""\
    JORDAN RIVERA
    jordan.rivera@example.com | (555) 010-0199 | Denver, CO
    PROFESSIONAL SUMMARY
    Analytics leader with 11 years of experience.
    EXPERIENCE
    Northwind Traders | Denver, CO Jan 2019 - Present
    Director, Analytics
    • Built a customer churn model that cut churn 12% in one
    year across every region.
    • Lead a team of 7 analysts.
    Senior Data Analyst 2016 - 2019
    Contoso Retail | Denver, CO
    • Automated weekly sales reporting, saving 10 hours per week.
    EDUCATION
    State University, Boulder, CO 2014
    B.S. Statistics
    SKILLS
    Python, SQL, Tableau
    """)


def test_heuristic_parser_handles_pdf_text():
    s = resume_ingest.verify(resume_ingest.heuristic_parse(PDF_LIKE), PDF_LIKE)
    assert [j.company for j in s.jobs] == ["Northwind Traders", "Contoso Retail"]
    nw, contoso = s.jobs
    assert (nw.title, nw.dates, nw.location) == ("Director, Analytics", "Jan 2019 - Present", "Denver, CO")
    assert nw.bullets[0] == "Built a customer churn model that cut churn 12% in one year across every region."
    assert contoso.title == "Senior Data Analyst", "title-first entries are swapped back"
    assert s.education[0].institution == "State University" and s.education[0].degree == "B.S. Statistics"
    assert s.summary.startswith("Analytics leader")
    assert not s.flags
    prefill = resume_ingest.prefill_contact(s)
    assert (prefill["name"], prefill["email"]) == ("Jordan Rivera", "jordan.rivera@example.com")
    assert prefill["skills"] == ["Python", "SQL", "Tableau"]


def test_verify_flags_text_not_in_source():
    s = resume_ingest.heuristic_parse(PDF_LIKE)
    s.jobs[0].bullets.append("Invented a time machine.")
    s.jobs[0].title = "Chief Everything Officer"
    flagged = {f.path for f in resume_ingest.verify(s, PDF_LIKE).flags}
    assert {"jobs.0.title", "jobs.0.bullets.2"} <= flagged


def test_structure_round_trips_through_pipeline_parser(tmp_path):
    s = resume_ingest.heuristic_parse(PDF_LIKE)
    s.jobs[0].company = "Northwind | Traders"  # a pipe would break the header line
    s.other_sections.append(resume_ingest.OtherSection(title="Certifications", lines=["Not | an employer"]))
    path = tmp_path / "in_resume_x.docx"
    resume_ingest.write_resume_docx(s, path)
    parsed = g.parse_prior_resume(path)
    assert [j["company"] for j in parsed["jobs"]] == ["Northwind / Traders", "Contoso Retail"]
    assert parsed["jobs"][0]["source_bullets"] == s.jobs[0].bullets
    back = resume_ingest.structure_from_paras(g.docx_paras(path, strip=False))
    assert back.summary == s.summary
    assert {o.title for o in back.other_sections} == {"Skills", "Certifications"}


def test_validate_structure_rejects_ambiguous_employers():
    s = resume_ingest.heuristic_parse(PDF_LIKE)
    s.jobs[1].company = s.jobs[0].company
    with pytest.raises(resume_ingest.IngestError, match="share the same company"):
        resume_ingest.validate_structure(s)


def test_ingest_structured_docx_parses_directly(tmp_path):
    make_examples.main(tmp_path)
    s = resume_ingest.ingest(tmp_path / "in_resume_example.docx")
    assert s.source == "docx"
    assert [j.company for j in s.jobs] == ["Northwind Traders", "Contoso Retail"]


# ---------------------------------------------------------------------------
# Templates and the render-only pipeline mode
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("template", sorted(g.TEMPLATES))
def test_templates_are_ats_safe_and_complete(tmp_path, template):
    out = tmp_path / f"{template}.docx"
    g.build_docx(g.result_to_docx_data(demo_data.DEMO_RESULT), out, template)
    d = docx.Document(str(out))
    text = "\n".join(p.text for p in d.paragraphs)
    assert not d.tables, "single column, no tables"
    assert not any(p.text.strip() for p in d.sections[0].header.paragraphs), "nothing in the header layer"
    for job in demo_data.DEMO_RESULT["experience"]:
        assert job["company"] in text and job["title"] in text
        for b in job["bullets"]:
            assert b["text"] in text
    assert "Jordan Rivera" in text


def _render(workspace, result_path, out_dir, *extra):
    return subprocess.run(
        [sys.executable, str(ROOT / "generate_resume.py"),
         "--profile", str(workspace.latest("in_profile")),
         "--resume", str(workspace.latest("in_resume")),
         "--job", str(ROOT / "examples" / "in_job_example.txt"),
         "--fact-bank", str(workspace.fact_bank_path),
         "--deny", str(workspace.deny_path),
         "--render-json", str(result_path), "--result-json", str(result_path),
         "--out-dir", str(out_dir), "--no-pdf", *extra],
        capture_output=True, text=True, encoding="utf-8", cwd=str(ROOT),
        env={**__import__("os").environ, "TAYLOR_PROGRESS": "1", "PYTHONIOENCODING": "utf-8"},
    )


def test_render_json_revalidates_and_writes(workspace, tmp_path):
    result = json.loads(json.dumps(demo_data.DEMO_RESULT))
    result["summary"] += " Also tried 97 things — all of them."
    result["experience"][0]["bullets"].append({"text": "Uncited but true: led a team of 7 analysts.", "ids": []})
    rp = tmp_path / "result.json"
    rp.write_text(json.dumps(result), encoding="utf-8")
    out = tmp_path / "out"
    proc = _render(workspace, rp, out, "--template", "modern")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "::stage rendering" in proc.stdout
    names = sorted(p.name for p in out.iterdir())
    assert any(n.endswith(".docx") for n in names) and any(n.endswith("_report.md") for n in names)
    saved = json.loads(rp.read_text(encoding="utf-8"))
    assert "—" not in saved["summary"], "em dashes are replaced on render"
    assert saved["template"] == "modern"
    assert any("97" in w for w in saved["warnings"])
    assert any("uncited bullet" in w for w in saved["warnings"])
    assert saved["coverage"]["covered"], "coverage is recomputed against the posting"

    # A second render archives the first one's outputs.
    archive = tmp_path / "archive"
    assert _render(workspace, rp, out, "--archive-dir", str(archive)).returncode == 0
    assert any(p.suffix == ".docx" for p in archive.iterdir())


def test_render_json_blocks_never_claim_hits(workspace, tmp_path):
    result = json.loads(json.dumps(demo_data.DEMO_RESULT))
    result["experience"][0]["bullets"][0]["text"] = "AWS certified solutions architect."
    rp = tmp_path / "result.json"
    rp.write_text(json.dumps(result), encoding="utf-8")
    out = tmp_path / "out"
    proc = _render(workspace, rp, out)
    assert proc.returncode != 0
    assert "::deny " in proc.stdout
    assert not out.exists() or not any(out.iterdir()), "nothing is written when a claim is prohibited"


# ---------------------------------------------------------------------------
# Project store
# ---------------------------------------------------------------------------


def test_project_store_lifecycle(tmp_path):
    store = ProjectStore(tmp_path / "projects", tmp_path / "projects" / "_trash")
    with pytest.raises(ProjectError):
        store.create(job_text="too short")
    meta = store.create(job_text=JOB_TEXT, company="Example Co.", role="Director, Customer Analytics")
    assert meta.name == "Director, Customer Analytics at Example Co."
    assert store.get_job_text(meta.id).startswith("Director")
    assert [m.id for m in store.all()] == [meta.id]

    store.update(meta.id, {"status": "applied", "id": "hijack", "notes": " call recruiter "})
    got = store.get(meta.id)
    assert (got.status, got.id, got.notes) == ("applied", meta.id, "call recruiter")
    with pytest.raises(ProjectError):
        store.update(meta.id, {"status": "hired?"})

    # Generation installs outputs; the next one moves them into versions/.
    for n in (1, 2):
        staging = store.dir(meta.id) / f".staging-{n}"
        staging.mkdir()
        (staging / "out_resume_x.docx").write_text(f"v{n}")
        (staging / "out_resume_x.pdf").write_text(f"v{n}")
        (staging / "result.json").write_text(json.dumps({**demo_data.DEMO_RESULT, "summary": f"v{n}"}))
        store.promote(meta.id, staging, new_version=True)
        time.sleep(1.05)  # versions are keyed by second
    assert store.get_result(meta.id)["summary"] == "v2"
    versions = store.versions(meta.id)
    assert len(versions) == 1 and versions[0]["summary"] == "v1"
    assert [f["name"] for f in versions[0]["files"]] == ["out_resume_x.docx", "result.json"], "PDFs aren't archived"

    with pytest.raises(NotFound):
        store.safe_file(meta.id, "outputs", "..", "project.json")
    with pytest.raises(NotFound):
        store.dir("../../etc")

    copy = store.duplicate(meta.id)
    assert copy.id != meta.id and store.get_result(copy.id)["summary"] == "v2"

    trash_id = store.trash_project(meta.id)
    assert meta.id not in [m.id for m in store.all()]
    assert store.list_trash()[0]["id"] == trash_id
    restored = store.restore(trash_id)
    assert restored.id in [m.id for m in store.all()]


def test_guess_company_role():
    assert guess_company_role(JOB_TEXT) == {"company": "Example Co.", "role": "Director, Customer Analytics"}
    assert guess_company_role("Title: Staff Analyst\nCompany: Contoso\n" + "x" * 50) == {"company": "Contoso", "role": "Staff Analyst"}


# ---------------------------------------------------------------------------
# Evidence: fact bank and never-claim list
# ---------------------------------------------------------------------------


def test_fact_bank_ids_are_stable_and_header_kept(workspace):
    facts = facts_mod.read_facts(workspace.fact_bank_path)
    facts.append(facts_mod.Fact(text="A brand new fact.", employer="general"))
    facts.append(facts_mod.Fact(id="F001", text="Duplicate id gets a fresh one."))
    facts_mod.write_facts(workspace.fact_bank_path, facts)
    raw = workspace.fact_bank_path.read_text(encoding="utf-8")
    assert raw.startswith("# Demo fact bank"), "the comment header survives a save"
    again = facts_mod.read_facts(workspace.fact_bank_path)
    ids = [f.id for f in again]
    assert ids[:11] == [f"F{n:03d}" for n in range(1, 12)]
    assert ids[11:] == ["F012", "F013"] and len(set(ids)) == len(ids)
    issues = facts_mod.employer_issues(
        [facts_mod.Fact(id="F900", employer="Nowhere Inc", text="x")], ["Northwind Traders"]
    )
    assert issues == [{"id": "F900", "employer": "Nowhere Inc"}]


def test_deny_list_validation(tmp_path):
    text = "# comment\n\\bcertified\\b\n(unclosed\n"
    checked = facts_mod.check_deny_text(text)
    assert [c["error"] is None for c in checked] == [True, False]
    with pytest.raises(facts_mod.FactBankError, match="Line 3"):
        facts_mod.write_deny_text(tmp_path / "deny.txt", text)
    assert facts_mod.match_phrase("\\bcertified\\b", "An AWS Certified pro") == ["\\bcertified\\b"]


def test_lint_flags_each_problem_on_its_line(workspace):
    result = json.loads(json.dumps(demo_data.DEMO_RESULT))
    result["experience"][0]["bullets"][0]["text"] = "Certified wizard, 42% better."
    result["experience"][0]["bullets"][1]["ids"] = ["F999"]
    result["experience"][1]["bullets"][0]["text"] += " Also at Northwind Traders."
    source = g.read_docx_text(workspace.latest("in_profile")) + g.read_docx_text(workspace.latest("in_resume"))
    bank = {f.id: f.model_dump() for f in facts_mod.read_facts(workspace.fact_bank_path)}
    out = lint_result(result, source, bank, g.load_deny_patterns(workspace.deny_path))
    kinds = lambda key: {i["kind"] for i in out["bullets"].get(key, [])}  # noqa: E731
    assert {"deny", "number"} <= kinds("0.0")
    assert "citation" in kinds("0.1")
    assert "cross_employer" in kinds("1.0")
    assert out["counts"]["error"] >= 2


# ---------------------------------------------------------------------------
# Jobs
# ---------------------------------------------------------------------------


def _wait(job, timeout=30):
    end = time.time() + timeout
    while not job.done and time.time() < end:
        time.sleep(0.05)
    assert job.done, "job did not finish"


def test_job_runner_parses_markers_and_hints(tmp_path):
    script = tmp_path / "fake.py"
    script.write_text(textwrap.dedent("""\
        import sys
        print("::stage drafting", flush=True)
        print("working...", flush=True)
        print('::deny {"where": "summary", "pattern": "x", "text": "y"}', flush=True)
        sys.exit("Could not find the `claude` CLI on PATH.")
        """))
    jobs = JobManager()
    finished = []
    job = jobs.submit("generate", "Test", [str(script)], on_finish=lambda j: finished.append(j.outcome))
    _wait(job)
    assert job.status == "failed" and finished == ["failed"]
    assert "drafting" in job.stages_seen and job.deny[0]["where"] == "summary"
    assert "working..." in job.lines and not any(l.startswith("::") for l in job.lines)
    assert "npm install -g @anthropic-ai/claude-code" in job.hint


def test_job_success_waits_for_on_success(tmp_path):
    script = tmp_path / "ok.py"
    script.write_text("print('::stage done')\n")
    jobs = JobManager()
    seen = []

    def on_success(j):
        seen.append(j.status)  # still "running" while results are installed
        time.sleep(0.2)

    job = jobs.submit("render", "Test", [str(script)], on_success=on_success)
    _wait(job)
    assert seen == ["running"] and job.status == "succeeded" and job.stage == "done"


# ---------------------------------------------------------------------------
# HTTP API
# ---------------------------------------------------------------------------

FAKE_PIPELINE = textwrap.dedent("""\
    # Stands in for generate_resume.py: writes a result and a .docx.
    import json, sys
    from pathlib import Path
    args = sys.argv[1:]
    get = lambda flag: args[args.index(flag) + 1] if flag in args else None
    out = Path(get("--out-dir")); out.mkdir(parents=True, exist_ok=True)
    for stage in ("preparing", "drafting", "structuring", "validating", "rendering"):
        print(f"::stage {stage}", flush=True)
    result = json.loads(Path(get("--render-json")).read_text()) if get("--render-json") else json.loads(RESULT)
    result["template"] = get("--template")
    Path(get("--result-json")).write_text(json.dumps(result))
    (out / "out_resume_jordan-rivera_2026-09-29.docx").write_text("docx")
    (out / "out_resume_jordan-rivera_2026-09-29_report.md").write_text("# Report")
    print("::stage done", flush=True)
    """)


@pytest.fixture
def client(workspace, tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from taylor.api import create_app

    fake = tmp_path / "fake_generate.py"
    fake.write_text(f"RESULT = {json.dumps(json.dumps(demo_data.DEMO_RESULT))}\n" + FAKE_PIPELINE, encoding="utf-8")
    monkeypatch.setattr(Paths, "generator_script", property(lambda self: fake))
    app = create_app(workspace, "tok", extra_hosts={"testserver"})
    c = TestClient(app)
    c.headers.update({"X-Taylor-Token": "tok"})
    return c


def test_api_requires_token_and_loopback_host(client):
    from fastapi.testclient import TestClient

    assert client.get("/api/health").status_code == 200
    bare = TestClient(client.app)
    assert bare.get("/api/projects").status_code == 403
    assert bare.get("/api/projects?t=tok").status_code == 200
    evil = TestClient(client.app, base_url="http://attacker.example")
    assert evil.get("/api/health", headers={"X-Taylor-Token": "tok"}).status_code == 421


def test_api_profile_and_resume_flow(client, workspace):
    body = client.get("/api/profile").json()
    assert body["exists"] and body["profile"]["contact"]["name"] == "Jordan Rivera"
    body["profile"]["sections"].append({"title": "Certifications", "items": ["None yet, and that's fine"]})
    saved = client.put("/api/profile", json=body["profile"]).json()
    assert saved["backup"] and saved["profile"]["sections"][-1]["title"] == "Certifications"

    bad = {**body["profile"], "contact": {**body["profile"]["contact"], "email": "nope"}}
    assert client.put("/api/profile", json=bad).status_code == 400

    with open(ROOT / "examples" / "in_resume_example.docx", "rb") as fh:
        up = client.post("/api/resume/upload", files={"file": ("my resume.docx", fh)}).json()
    assert not up["needs_ai"] and up["structure"]["jobs"][0]["company"] == "Northwind Traders"
    confirmed = client.post("/api/resume/confirm", json={"structure": up["structure"], "upload_id": up["upload_id"]}).json()
    assert confirmed["file"] == "in_resume_Jordan-Rivera.docx"
    # The demo fact bank names Fabrikam Health, which this smaller resume doesn't have.
    assert any(i["employer"] == "Fabrikam Health" for i in confirmed["fact_bank_issues"])
    assert client.post("/api/resume/upload", files={"file": ("x.exe", b"MZ")}).status_code == 400


def test_api_project_generate_edit_render(client):
    created = client.post("/api/projects", json={"job_text": JOB_TEXT, "company": "Example Co.", "role": "Director"}).json()
    pid = created["project"]["id"]
    job = client.post(f"/api/projects/{pid}/generate").json()

    events = []
    with client.stream("GET", f"/api/jobs/{job['id']}/events") as resp:
        for line in resp.iter_lines():
            if line.startswith("event:"):
                events.append(line.split(":", 1)[1].strip())
            if events and events[-1] == "end" and line.startswith("data:"):
                final = json.loads(line[5:])
                break
    assert events[0] in ("lines", "status") and events[-1] == "end"
    assert final["status"] == "succeeded"

    detail = client.get(f"/api/projects/{pid}").json()
    assert detail["result"]["summary"] == demo_data.DEMO_RESULT["summary"]
    assert detail["project"]["last_generate"]["ok"] and not detail["project"]["inputs_changed"]
    assert [f["name"] for f in detail["outputs"]][0].startswith("out_resume_")

    result = detail["result"]
    result["summary"] = "Edited by hand."
    assert client.put(f"/api/projects/{pid}/result", json=result).json()["outputs_stale"] is True
    lint = client.post(f"/api/projects/{pid}/lint", json=result).json()
    assert "counts" in lint

    job = client.post(f"/api/projects/{pid}/render", json={"template": "compact"}).json()
    end = time.time() + 20
    while client.get(f"/api/jobs/{job['id']}").json()["status"] not in ("succeeded", "failed") and time.time() < end:
        time.sleep(0.05)
    detail = client.get(f"/api/projects/{pid}").json()
    assert detail["project"]["template"] == "compact" and detail["result"]["summary"] == "Edited by hand."

    name = detail["outputs"][0]["name"]
    assert client.get(f"/api/projects/{pid}/files/{name}").status_code == 200
    assert client.get(f"/api/projects/{pid}/files/..%2Fproject.json").status_code == 404
    assert client.get(f"/api/projects/{pid}/report").text == "# Report"

    trash = client.delete(f"/api/projects/{pid}").json()["trash_id"]
    assert client.get(f"/api/projects/{pid}").status_code == 404
    assert client.post(f"/api/trash/{trash}/restore").status_code == 200


def test_api_import_legacy_postings(client, workspace):
    (workspace.input_dir / "in_job_example.txt").write_text(JOB_TEXT, encoding="utf-8")
    candidates = client.get("/api/import/candidates").json()
    assert candidates == [{"file": "in_job_example.txt", "company": "Example Co.", "role": "Director, Customer Analytics"}]
    assert len(client.post("/api/import", json={"files": ["in_job_example.txt"]}).json()["created"]) == 1
    assert client.get("/api/import/candidates").json() == []


# ---------------------------------------------------------------------------
# Portability: API key, sample data, data root
# ---------------------------------------------------------------------------


FAKE_KEY = "sk-ant-api03-" + "A1b2C3d4E5" * 3


def _client(paths):
    from fastapi.testclient import TestClient

    app = create_app(paths, "tok", extra_hosts={"testserver"})
    return TestClient(app, headers={"X-Taylor-Token": "tok"})


def test_api_key_is_never_returned_or_settings_leaked(workspace):
    c = _client(workspace)
    assert c.put("/api/settings/api-key", json={"key": "not a key"}).status_code == 400
    r = c.put("/api/settings/api-key", json={"key": FAKE_KEY})
    assert r.json() == {"set": True, "source": "app", "last4": FAKE_KEY[-4:]}
    for url in ("/api/settings", "/api/system", "/api/settings/api-key"):
        assert FAKE_KEY not in c.get(url).text
    assert credentials_mod.load_key(workspace.secrets_path) == FAKE_KEY
    assert c.delete("/api/settings/api-key").json()["set"] is False
    assert not workspace.secrets_path.exists()


def test_api_key_reaches_jobs_and_logs_are_redacted(workspace):
    assert credentials_mod.redact(f"boom {FAKE_KEY} end") == "boom sk-ant-*** end"


def test_generator_env_keeps_only_the_api_key(monkeypatch):
    import generate_resume as g

    monkeypatch.setenv("ANTHROPIC_API_KEY", FAKE_KEY)
    monkeypatch.setenv("ANTHROPIC_BASE_URL", "https://example.invalid")
    monkeypatch.setenv("CLAUDE_CODE_FOO", "1")
    env = g._clean_subprocess_env()
    assert env.get("ANTHROPIC_API_KEY") == FAKE_KEY
    assert "ANTHROPIC_BASE_URL" not in env and "CLAUDE_CODE_FOO" not in env


def test_default_data_root_is_outside_the_repo(monkeypatch, tmp_path):
    from taylor.paths import default_data_root

    monkeypatch.delenv("RESUME_TAYLOR_DATA", raising=False)
    monkeypatch.delenv("RESUME_STUDIO_DATA", raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    assert default_data_root() == tmp_path / "ResumeTaylor"
    monkeypatch.setenv("RESUME_TAYLOR_DATA", str(tmp_path / "mine"))
    assert default_data_root() == (tmp_path / "mine").resolve()


def test_default_data_root_adopts_the_old_resume_studio_folder(monkeypatch, tmp_path):
    from taylor.paths import default_data_root

    monkeypatch.delenv("RESUME_TAYLOR_DATA", raising=False)
    monkeypatch.delenv("RESUME_STUDIO_DATA", raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    (tmp_path / "ResumeStudio" / "projects").mkdir(parents=True)
    (tmp_path / "ResumeStudio" / "secrets.json").write_text("{}", encoding="utf-8")

    root = default_data_root()
    assert root == tmp_path / "ResumeTaylor"
    assert (root / "secrets.json").read_text(encoding="utf-8") == "{}"
    assert (root / "projects").is_dir() and not (tmp_path / "ResumeStudio").exists()

    # An existing new folder is never merged into or overwritten.
    (tmp_path / "ResumeStudio").mkdir()
    assert default_data_root() == tmp_path / "ResumeTaylor"
    assert (tmp_path / "ResumeStudio").exists()

    monkeypatch.setenv("RESUME_STUDIO_DATA", str(tmp_path / "old-override"))
    assert default_data_root() == (tmp_path / "old-override").resolve()


def test_sample_data_loads_only_into_an_empty_folder_and_clears(tmp_path):
    paths = Paths(tmp_path)
    paths.ensure()
    c = _client(paths)
    assert c.get("/api/system").json()["onboarding_needed"] is True
    assert c.post("/api/demo/load").status_code == 200
    system = c.get("/api/system").json()
    assert system["sample_loaded"] and not system["onboarding_needed"]
    assert c.post("/api/demo/load").status_code == 409
    assert c.post("/api/demo/clear").json()["backup"] is None
    system = c.get("/api/system").json()
    assert not system["sample_loaded"] and system["onboarding_needed"]
    assert not [p for p in paths.projects_dir.iterdir() if p.is_dir() and not p.name.startswith("_")]


def test_sample_data_refuses_to_mix_with_a_real_profile(workspace):
    assert _client(workspace).post("/api/demo/load").status_code == 409
