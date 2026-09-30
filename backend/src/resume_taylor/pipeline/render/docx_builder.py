"""Building the resume and cover-letter .docx files from a template spec."""

from datetime import datetime
from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .templates import DEFAULT_TEMPLATE, RULE_COLOR, TEMPLATES, TemplateSpec, get_template


def _set_bottom_border(paragraph, color: str = RULE_COLOR, size: int = 6) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    """Append a clickable external link (python-docx has no API for this).
    Styled like the surrounding contact text so it doesn't read as blue link
    text; ATS parsers and PDF viewers still pick up the target."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    run_text = OxmlElement("w:t")
    run_text.text = text
    run_text.set(qn("xml:space"), "preserve")
    run.append(run_text)
    link.append(run)
    paragraph._p.append(link)


def _build_styles(d: docx.Document, spec: TemplateSpec) -> dict:
    styles = d.styles

    normal = styles["Normal"]
    normal.font.name = spec.font
    normal.font.size = Pt(spec.body_size)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(spec.body_space_after)
    normal.paragraph_format.line_spacing = 1.0

    align = WD_ALIGN_PARAGRAPH.CENTER if spec.name_align == "center" else WD_ALIGN_PARAGRAPH.LEFT

    name_style = styles.add_style("ResumeName", WD_STYLE_TYPE.PARAGRAPH)
    name_style.base_style = normal
    name_style.font.size = Pt(spec.name_size)
    name_style.font.bold = True
    name_style.font.small_caps = spec.name_small_caps or None
    name_style.font.color.rgb = spec.ink
    name_style.paragraph_format.alignment = align
    name_style.paragraph_format.space_after = Pt(2)

    contact_style = styles.add_style("ResumeContact", WD_STYLE_TYPE.PARAGRAPH)
    contact_style.base_style = normal
    contact_style.font.size = Pt(spec.contact_size)
    contact_style.font.color.rgb = spec.muted
    contact_style.paragraph_format.alignment = align
    contact_style.paragraph_format.space_after = Pt(spec.contact_space_after)

    section_style = styles.add_style("ResumeSection", WD_STYLE_TYPE.PARAGRAPH)
    section_style.base_style = normal
    section_style.font.size = Pt(spec.section_size)
    section_style.font.bold = True
    section_style.font.color.rgb = spec.ink
    section_style.font.all_caps = spec.section_all_caps or None
    section_style.font.small_caps = spec.section_small_caps or None
    section_style.paragraph_format.space_before = Pt(spec.section_space_before)
    section_style.paragraph_format.space_after = Pt(spec.section_space_after)

    entry_title_style = styles.add_style("ResumeEntryTitle", WD_STYLE_TYPE.PARAGRAPH)
    entry_title_style.base_style = normal
    entry_title_style.font.italic = True
    entry_title_style.font.size = Pt(spec.entry_title_size)
    entry_title_style.font.color.rgb = spec.muted
    entry_title_style.paragraph_format.space_after = Pt(3 if spec.body_space_after >= 4 else 1.5)

    bullet_style = styles["List Bullet"]
    bullet_style.font.name = spec.font
    bullet_style.font.size = Pt(spec.body_size)  # match body text in Summary/Education/Skills
    bullet_style.paragraph_format.left_indent = Inches(spec.bullet_indent)
    bullet_style.paragraph_format.space_after = Pt(spec.bullet_space_after)
    bullet_style.paragraph_format.line_spacing = 1.0

    return {"section": section_style, "entry_title": entry_title_style}


def _add_entry_header(d: docx.Document, left_text: str, right_text: str, spec: TemplateSpec | None = None):
    spec = spec or TEMPLATES[DEFAULT_TEMPLATE]
    p = d.add_paragraph()
    content_width = (
        d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin
    )
    p.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
    left_run = p.add_run(left_text)
    left_run.bold = True
    if right_text:
        right_run = p.add_run(f"\t{right_text}")
        right_run.font.color.rgb = spec.muted
    p.paragraph_format.space_after = Pt(1)
    return p


def _new_document(data: dict, spec: TemplateSpec | None = None) -> tuple:
    """Letter-size document with the shared styles and the name/contact
    header, used by both the resume and the cover letter."""
    spec = spec or TEMPLATES[DEFAULT_TEMPLATE]
    d = docx.Document()

    section = d.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(spec.margin_tb)
    section.bottom_margin = Inches(spec.margin_tb)
    section.left_margin = Inches(spec.margin_lr)
    section.right_margin = Inches(spec.margin_lr)

    styles = _build_styles(d, spec)

    d.add_paragraph(f"{data['first_name']} {data['last_name']}", style="ResumeName")
    contact_p = d.add_paragraph(style="ResumeContact")
    linkedin = data.get("linkedin", "")
    parts = [
        (data["email"], f"mailto:{data['email']}" if data["email"] else None),
        (data["phone"], None),
        (data["location"], None),
        (linkedin, linkedin if linkedin.lower().startswith("http") else f"https://{linkedin}"),
    ]
    first = True
    for text, url in parts:
        if not text:
            continue
        if not first:
            contact_p.add_run(" | ")
        first = False
        if url:
            _add_hyperlink(contact_p, text, url)
        else:
            contact_p.add_run(text)
    return d, styles


def build_cover_letter_docx(data: dict, paragraphs: list, out_path: Path, template: str | None = None) -> None:
    d, _styles = _new_document(data, get_template(template))
    d.add_paragraph(datetime.now().strftime("%B %d, %Y").replace(" 0", " "))
    d.add_paragraph("Dear Hiring Team,")
    for text in paragraphs:
        d.add_paragraph(text).paragraph_format.space_after = Pt(8)
    d.add_paragraph("Sincerely,")
    d.add_paragraph(f"{data['first_name']} {data['last_name']}")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))


def build_docx(data: dict, out_path: Path, template: str | None = None) -> None:
    spec = get_template(template)
    d, styles = _new_document(data, spec)

    def add_heading(text: str) -> None:
        p = d.add_paragraph(text, style=styles["section"])
        if spec.rule_color:
            _set_bottom_border(p, spec.rule_color, spec.rule_size)

    add_heading("Summary")
    d.add_paragraph(data["summary"])

    if data.get("education"):
        add_heading("Education")
        for edu in data["education"]:
            header = " | ".join(
                part for part in (edu.get("institution"), edu.get("location")) if part
            )
            _add_entry_header(d, header, edu.get("dates", ""), spec)
            degree_p = d.add_paragraph(edu.get("degree", ""))
            degree_p.paragraph_format.space_after = Pt(6 if spec.body_space_after >= 4 else 3)

    if data.get("experience"):
        add_heading("Experience")
        for job in data["experience"]:
            place = " | ".join(part for part in (job.get("company"), job.get("location")) if part)
            if spec.entry_layout == "title_first":
                _add_entry_header(d, job.get("title", "") or place, job.get("dates", ""), spec)
                company_p = d.add_paragraph()
                company_run = company_p.add_run(place)
                company_run.font.color.rgb = spec.accent
                company_run.bold = True
                company_run.font.size = Pt(spec.entry_title_size)
                company_p.paragraph_format.space_after = Pt(3)
            else:
                _add_entry_header(d, place, job.get("dates", ""), spec)
                d.add_paragraph(job.get("title", ""), style=styles["entry_title"])
            bullets = job.get("bullets", [])
            for i, bullet in enumerate(bullets):
                p = d.add_paragraph(bullet, style="List Bullet")
                if i == len(bullets) - 1:
                    p.paragraph_format.space_after = Pt(spec.job_space_after)

    if data.get("skills"):
        add_heading("Skills")
        d.add_paragraph(spec.skills_separator.join(data["skills"]))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))
