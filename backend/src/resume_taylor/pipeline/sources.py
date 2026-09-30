"""Reading the pipeline's input files: .docx / .pdf / .txt / .md text, the newest
`in_*` file by modification time, and the (style, text) paragraph lists the parsers work on."""

import sys
from pathlib import Path

import docx

from ..config import INPUT_DIR


def open_docx(path: Path) -> docx.Document:
    # python-docx reports a file Word has open and locked as "Package not
    # found", which hides the real (and easily fixed) cause.
    try:
        with open(path, "rb"):
            pass
    except PermissionError:
        sys.exit(f"{path.name} is locked, most likely open in Word. Close it and re-run.")
    return docx.Document(str(path))


def read_docx_text(path: Path) -> str:
    d = open_docx(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


JOB_POSTING_EXTS = (".docx", ".pdf", ".txt", ".md")


def read_input_text(path: Path) -> str:
    """Plain text of an input whose structure isn't parsed (the job posting),
    so it can be saved however is convenient: .docx, .pdf, .txt, or .md."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages).strip()
    return path.read_text(encoding="utf-8", errors="replace").strip()


def find_latest(prefix: str, exts: tuple = (".docx",), directory: Path | None = None) -> Path:
    """Newest matching file by modification time (not name, which would pick
    an alphabetically-last posting over the one just dropped in). Looks in
    resume_input/ unless `directory` is given."""
    directory = directory or INPUT_DIR
    candidates = [p for p in directory.glob(f"{prefix}*") if p.suffix.lower() in exts]
    matches = sorted(candidates, key=lambda p: p.stat().st_mtime)
    if not matches:
        sys.exit(
            f"No {'/'.join(exts)} file starting with '{prefix}' found in {directory}. "
            f"Expected a profile and prior resume (in_profile*.docx, in_resume*.docx) and a "
            f"job posting (in_job*: {', '.join(JOB_POSTING_EXTS)})."
        )
    return matches[-1]


def docx_paras(path: Path, strip: bool = True) -> list:
    """The (style name, text) pairs of every non-empty paragraph. The parsers
    below work on this list, so a caller that already has the structure (the
    Resume Taylor profile editor, a normalized uploaded resume) can hand it
    over directly instead of round-tripping through a file."""
    doc = open_docx(path)
    return [
        (p.style.name, p.text.strip() if strip else p.text) for p in doc.paragraphs if p.text.strip()
    ]


def paras_and_name(source, strip: bool = True) -> tuple:
    if isinstance(source, (str, Path)):
        return docx_paras(Path(source), strip), Path(source).name
    return [(style, text.strip() if strip else text) for style, text in source], "the profile"
