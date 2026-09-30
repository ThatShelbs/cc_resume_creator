"""The .docx builders and resume templates."""

import docx
import pytest

from resume_taylor.pipeline.render.docx_builder import build_cover_letter_docx, build_docx
from resume_taylor.pipeline.render.templates import TEMPLATES
from resume_taylor.pipeline.result import result_to_docx_data
from resume_taylor.pipeline.sources import read_docx_text
from resume_taylor.sample_data import demo_data


def test_build_docx_contains_content(tmp_path):
    data = {
        "first_name": "Ada",
        "last_name": "Lovelace",
        "email": "ada@example.com",
        "phone": "",
        "location": "London",
        "summary": "Analyst.",
        "education": [],
        "experience": [
            {"company": "Acme", "location": "", "dates": "2020", "title": "Lead", "bullets": ["Did X."]}
        ],
        "skills": ["Python"],
    }
    out = tmp_path / "out.docx"
    build_docx(data, out)
    text = read_docx_text(out)
    for expected in ("Ada Lovelace", "Analyst.", "Did X.", "Python"):
        assert expected in text


def test_contact_links_are_hyperlinks(tmp_path):
    data = {
        "first_name": "Ada",
        "last_name": "Lovelace",
        "email": "ada@example.com",
        "phone": "(555) 010-0199",
        "location": "London",
        "linkedin": "linkedin.com/in/ada",
        "summary": "Analyst.",
        "skills": [],
    }
    out = tmp_path / "out.docx"
    build_docx(data, out)
    d = docx.Document(str(out))
    targets = {r.target_ref for r in d.part.rels.values() if r.is_external}
    assert targets == {"mailto:ada@example.com", "https://linkedin.com/in/ada"}
    assert "ada@example.com | (555) 010-0199 | London | linkedin.com/in/ada" in read_docx_text(out)


def test_build_cover_letter_docx(tmp_path):
    data = {"first_name": "Ada", "last_name": "Lovelace", "email": "ada@example.com", "phone": "", "location": ""}
    out = tmp_path / "cl.docx"
    build_cover_letter_docx(data, ["First paragraph.", "Second paragraph."], out)
    text = read_docx_text(out)
    for expected in ("Dear Hiring Team,", "First paragraph.", "Second paragraph.", "Sincerely,"):
        assert expected in text


@pytest.mark.parametrize("template", sorted(TEMPLATES))
def test_templates_are_ats_safe_and_complete(tmp_path, template):
    out = tmp_path / f"{template}.docx"
    build_docx(result_to_docx_data(demo_data.DEMO_RESULT), out, template)
    d = docx.Document(str(out))
    text = "\n".join(p.text for p in d.paragraphs)
    assert not d.tables, "single column, no tables"
    assert not any(p.text.strip() for p in d.sections[0].header.paragraphs), "nothing in the header layer"
    for job in demo_data.DEMO_RESULT["experience"]:
        assert job["company"] in text and job["title"] in text
        for b in job["bullets"]:
            assert b["text"] in text
    assert "Jordan Rivera" in text
